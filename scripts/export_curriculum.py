"""Export the included curriculum source to a Word reference document.

Run from the project root: python3 scripts/export_curriculum.py
Requires python-docx. This is an editorial DRAFT, not an approved syllabus.
"""
from __future__ import annotations

import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / 'src/curriculum/master.json').read_text())
OUTPUT = ROOT / 'Desloop_Academy_Master_Curriculum.docx'
NAVY = RGBColor(20, 35, 68)
GREEN = RGBColor(21, 105, 80)
GREY = RGBColor(74, 85, 103)
WHITE = RGBColor(255, 255, 255)
AGE_LABELS = [('early', 'Ages 6–8'), ('junior', 'Ages 9–12'), ('teen', 'Ages 13–17'), ('adult', 'Ages 18+')]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, colour=NAVY):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    run.font.color.rgb = colour
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_list(doc, entries):
    for entry in entries:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(entry)


def add_labelled(doc, label, content):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    lead = p.add_run(label + '  ')
    lead.bold = True
    lead.font.color.rgb = NAVY
    p.add_run(content)
    return p


def page_number(p):
    run = p.add_run()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.77)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(9.5)
normal.font.color.rgb = GREY
normal.paragraph_format.line_spacing = 1.14
normal.paragraph_format.space_after = Pt(6)
for name, size, colour in [('Title', 32, NAVY), ('Heading 1', 22, NAVY), ('Heading 2', 15, NAVY), ('Heading 3', 11, GREEN)]:
    style = styles[name]
    style.font.name = 'Aptos Display'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = colour
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(16 if name != 'Heading 3' else 8)
    style.paragraph_format.space_after = Pt(7)
styles['List Bullet'].paragraph_format.space_after = Pt(2)
styles['List Number'].paragraph_format.space_after = Pt(2)

header = section.header.paragraphs[0]
header.text = 'DESLOOP  /  ACADEMY                                      EDITORIAL CURRICULUM DRAFT'
header.style = 'Caption'
header.runs[0].font.color.rgb = GREEN
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Draft-1  •  21 October 2026  |  ')
page_number(footer)
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = GREY

doc.core_properties.title = 'Desloop Academy — Master Curriculum'
doc.core_properties.subject = '44-course sequential draft curriculum for review and implementation'
doc.core_properties.author = 'Desloop Academy — draft preview'
doc.core_properties.keywords = 'curriculum, 44 courses, progression, teaching bands, draft'

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(70)
r = p.add_run('DESLOOP ACADEMY')
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = GREEN
p = doc.add_paragraph(style='Title')
p.add_run('Master Curriculum\nOne pathway. Many ways to learn.')
p = doc.add_paragraph()
r = p.add_run('44 required courses  /  9 ordered stages  /  176 module lesson drafts')
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = NAVY
p.paragraph_format.space_before = Pt(14)
p = doc.add_paragraph('Course-level expansion, age-adapted teaching scaffolds, projects, assessment policy and implementation notes. Generated from the application’s included curriculum source so the reference and frontend stay aligned.')
p.paragraph_format.space_before = Pt(18)
p = doc.add_paragraph()
r = p.add_run('DRAFT FOR REVIEW — NOT A PUBLISHED OR ACCREDITED SYLLABUS')
r.bold = True
r.font.color.rgb = GREEN
p = doc.add_paragraph('Version draft-1  •  21 October 2026  •  Ghana-first, internationally adaptable')
p.paragraph_format.space_before = Pt(18)
doc.add_page_break()

doc.add_heading('Read this before using the draft', 1)
doc.add_paragraph('Every learner begins at Course 01 and moves through all 44 courses in the same order. Age bands alter examples, session length, support and response modes; they do not remove an essential outcome or open a later course early. Optional enrichment and remediation remain inside the active course.')
doc.add_paragraph('All objectives, examples and five-item checks are authoring drafts. The working frontend intentionally labels them as such. Subject experts, child-safety specialists, accessibility reviewers, teachers in each age band and legal reviewers must validate them before a real learner uses them. No qualification, job or certification is promised.')

doc.add_heading('Progression and pass policy', 2)
add_list(doc, [
    'Courses unlock only in sequence: Course N+1 requires a version-matched, valid pass for Course N and every earlier course.',
    'Complete all required module lessons and documented practice; the five-item course knowledge check averages at least 80% (four of five correct).',
    'Submit a project artefact or an accessible, supervised description of it. An authorised human reviewer assigns a project rubric of at least 70% and records feedback.',
    'All safety-critical items must be correct. A human safety/privacy check must pass; an unsafe project never passes on numerical score alone.',
    'A failed gate keeps the learner in the same course with targeted explanation, fresh practice, a retry and a documented review/revision state.',
    'Default overall demonstration weights: quiz 20%, guided practice 20%, project 40%, reflection 10%, safety 10%. The final capstone uses its own published seven-part rubric.',
    'The frontend demonstrates these gates with local browser state; official completion, authorisation, approval, audit, versioning and unlocking must happen on a secure backend.',
])

doc.add_heading('Age and accessibility without a second track', 2)
for label, text in [
    ('Ages 6–8', 'Short, picture-led, adult-supervised sessions; paper/simulation alternatives; oral or drawn evidence may be recorded by a facilitator. Never demand a public child profile.'),
    ('Ages 9–12', 'Concrete examples, visual checklists, approved tools and a modest independent step, with consent-aware collaboration.'),
    ('Ages 13–17', 'More reasoning, debugging and documented trade-offs with moderated sharing and privacy protections.'),
    ('Ages 18+', 'Applied study, workplace or community scenarios with fuller documentation and independent evaluation.'),
]:
    add_labelled(doc, label, text)
doc.add_paragraph('Accessibility choices such as zoom, captions, keyboard-only use, plain language, high contrast and assisted response formats must be available without changing the pass criteria. Avoid punitive speed thresholds. Use fictional or synthetic data, especially for children. Practical coding, robotics and cybersecurity activities require age-appropriate approved environments.')

doc.add_heading('At a glance', 2)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Light Shading Accent 1'
for cell, title in zip(table.rows[0].cells, ['Stage', 'Theme', 'Required courses']):
    shade(cell, '142344')
    set_cell_text(cell, title, True, WHITE)
for stage in DATA['stages']:
    cells = table.add_row().cells
    for cell, text in zip(cells, [stage['number'], stage['title'], f"{stage['first']:02d}–{stage['last']:02d}"]):
        set_cell_text(cell, text)

doc.add_heading('How this file relates to the system', 2)
add_labelled(doc, 'Single runtime source', 'src/curriculum/master.json (regenerate editorial seed with python3 scripts/seed_master.py). The Word export reads that JSON; it does not maintain a separate syllabus.')
add_labelled(doc, 'Learner routes', '/courses for the map; /courses/:slug for course outlines; /app/courses for progress; /app/learn/:courseId/:lessonId and /app/assess/:courseId for gated work.')
add_labelled(doc, 'Reviewer routes', '/facilitator/pathway-reviews and /classroom/pathway-reviews preview human decisions. Local role switching is not authorisation.')
doc.add_page_break()

for stage in DATA['stages']:
    doc.add_heading(f"Stage {stage['number']:02d}  |  {stage['title']}", 1)
    doc.add_paragraph(stage['subtitle'])
    doc.add_paragraph(f"Required courses {stage['first']:02d}–{stage['last']:02d}; no direct entry into a later stage.")
    for course in DATA['courses']:
        if not stage['first'] <= course['number'] <= stage['last']:
            continue
        doc.add_heading(f"Course {course['number']:02d}  —  {course['title']}", 2)
        add_labelled(doc, 'Purpose', course['purpose'])
        add_labelled(doc, 'Suggested guided time', f"~{course['guidedHours']} hours, split into age-appropriate sessions; not a timed pass condition.")
        add_labelled(doc, 'Editorial state', f"{course['version']} · unreviewed teaching and assessment drafts")
        doc.add_heading('Demonstrable outcomes', 3)
        add_list(doc, course['outcomes'])
        doc.add_heading('Required project and evidence', 3)
        doc.add_paragraph(course['project']['brief'])
        add_list(doc, course['project']['requiredEvidence'])
        add_labelled(doc, 'Assessment focus', course['assessment'])
        if course['safetyCritical']:
            add_labelled(doc, 'Critical gate', 'The safety/privacy check is mandatory. A missed critical quiz item or failed human safeguard assessment prevents a pass, even if the numeric average is high.')
        doc.add_heading('Four ordered teaching modules', 3)
        for number, module in enumerate(course['modules'], 1):
            lesson = module['lessons'][0]
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            r = p.add_run(f"{number}. {module['title']}  ·  {lesson['minutes']} min suggested segment")
            r.bold = True
            r.font.color.rgb = NAVY
            add_labelled(doc, 'Objective', lesson['objective'])
            add_labelled(doc, 'Guided example', lesson['workedExample'])
            add_labelled(doc, 'Required practice', lesson['independentPractice'])
            add_labelled(doc, 'Draft formative check', lesson['check']['question'] + ' Intended response: ' + lesson['check']['options'][lesson['check']['correct']])
        doc.add_heading('Age-adapted facilitation', 3)
        for key, label in AGE_LABELS:
            add_labelled(doc, label, course['ageAdaptation'][key])
        p = doc.add_paragraph('Progression: pass the prior course; complete these four lessons and practice items; meet the course quiz and safety gates; submit this project for a rubric and a human decision. If changes are needed, remain here and retry.')
        p.paragraph_format.space_before = Pt(6)
        p.runs[0].italic = True


doc.add_page_break()
doc.add_heading('Final capstone rubric', 1)
doc.add_paragraph('Course 44 integrates planning, making, testing, safety, accessibility, reflection and communication. Its weighted preview rubric is separate from the default course weighting. A score of at least 70% and a passing critical safety/privacy gate are required, along with the 80% knowledge check, required practice and human review.')
capstone = [
    ('Problem and scope', '10%'), ('Functionality', '25%'), ('Accessibility and design', '15%'),
    ('Technical reasoning', '15%'), ('Testing and revision', '15%'), ('Safety and privacy', '10%'),
    ('Presentation and reflection', '10%'),
]
for label, weight in capstone:
    add_labelled(doc, weight, label)

doc.add_heading('Before any real publication', 1)
add_list(doc, [
    'Replace generated lesson outlines with subject-expert-authored scripts and worked examples; validate every quiz distractor, correct answer, safety scenario and rubric anchor.',
    'Test the full pathway with teachers and learners in each age band. Validate reading level, translation/localisation, accessibility and realistic duration, including non-text response alternatives.',
    'Complete Ghana-first legal, privacy, child safeguarding, content rights and institutional review before real child accounts, uploads, discussion spaces or AI access.',
    'Make the backend the sole writer of attempts, rubric decisions, course-version assignments and unlock events. Verify reviewer identity, tenant membership and guardian permissions; preserve an audit trail.',
    'Plan migration when a published course or prerequisite policy changes; never silently discard a valid prior completion or let a stale client record create a pass.',
    'Run real moderation, incident escalation, content review, security testing, accessibility testing and offline-sync reconciliation before production launch.',
])

doc.save(OUTPUT)
print(f'Wrote {OUTPUT} ({OUTPUT.stat().st_size:,} bytes; {len(DATA["courses"])} courses)')
